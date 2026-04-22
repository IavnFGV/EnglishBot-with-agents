import asyncio
import sys
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from aiogram.types import Document, User
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from englishbot import db
from englishbot.topics import create_topic_for_teacher_workspace
from englishbot.user_profiles import set_user_role
from englishbot.vocabulary import (
    create_learning_item_for_teacher_workspace,
    create_learning_item_translation,
    create_lexeme,
)
from englishbot.workbook_export import LEARNING_ITEMS_COLUMNS, LEARNING_ITEMS_SHEET_NAME, META_COLUMNS, META_SHEET_NAME
from englishbot.workbook_handlers import workbook_export, workbook_import_document, workbook_import_usage
from englishbot.workspaces import add_workspace_member, create_workspace


class FakeBot:
    def __init__(self, download_payload: bytes | None = None) -> None:
        self.download_payload = download_payload or b""

    async def download(self, file: object, destination: BytesIO) -> BytesIO:
        destination.write(self.download_payload)
        destination.seek(0)
        return destination


class FakeMessage:
    def __init__(
        self,
        user: User,
        *,
        bot: FakeBot | None = None,
        caption: str | None = None,
        document: Document | None = None,
    ) -> None:
        self.from_user = user
        self.bot = bot or FakeBot()
        self.caption = caption
        self.document = document
        self.answers: list[dict[str, object]] = []
        self.documents: list[dict[str, object]] = []

    async def answer(self, text: str, **kwargs: object) -> None:
        self.answers.append({"text": text, "kwargs": kwargs})

    async def answer_document(self, document: object, **kwargs: object) -> None:
        self.documents.append({"document": document, "kwargs": kwargs})


def make_user(user_id: int, first_name: str) -> User:
    return User(id=user_id, is_bot=False, first_name=first_name, username=first_name.lower())


def make_document(file_name: str) -> Document:
    return Document(
        file_id="file-id",
        file_unique_id="file-unique-id",
        file_name=file_name,
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        file_size=123,
    )


def setup_db(tmp_path: Path) -> None:
    db.DB_PATH = tmp_path / "workbook_handlers.sqlite3"
    db.init_db()


def create_teacher_workspace(tmp_path: Path, teacher_user_id: int = 1001) -> tuple[User, int]:
    setup_db(tmp_path)
    teacher = make_user(teacher_user_id, "Teacher")
    db.save_user(teacher)
    set_user_role(teacher.id, "teacher")
    workspace = create_workspace("Authoring", kind="teacher")
    add_workspace_member(workspace["workspace_id"], teacher.id, "teacher")
    return teacher, int(workspace["workspace_id"])


def seed_workspace_content(teacher_user_id: int, workspace_id: int) -> None:
    lexeme_id = create_lexeme("cat")
    learning_item_id = create_learning_item_for_teacher_workspace(
        teacher_user_id,
        workspace_id,
        lexeme_id,
        "cat",
    )
    create_learning_item_translation(learning_item_id, "ru", "кот")
    create_topic_for_teacher_workspace(teacher_user_id, workspace_id, "animals", "Animals")


def build_workbook_bytes() -> bytes:
    workbook = Workbook()
    meta_sheet = workbook.active
    meta_sheet.title = META_SHEET_NAME
    meta_sheet.append(META_COLUMNS)
    meta_sheet.append(("2", "Authoring", "ru"))

    learning_items_sheet = workbook.create_sheet(LEARNING_ITEMS_SHEET_NAME)
    learning_items_sheet.append(LEARNING_ITEMS_COLUMNS)
    learning_items_sheet.append(
        (
            "cat",
            "кот",
            None,
            "animals",
            None,
            None,
            None,
            None,
            0,
        )
    )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_workbook_export_handler_returns_xlsx_document(tmp_path: Path) -> None:
    teacher, workspace_id = create_teacher_workspace(tmp_path)
    seed_workspace_content(teacher.id, workspace_id)
    message = FakeMessage(teacher)

    asyncio.run(workbook_export(message, SimpleNamespace(args=str(workspace_id))))

    assert message.answers == []
    assert len(message.documents) == 1
    sent_document = message.documents[0]["document"]
    assert sent_document.filename == f"teacher-workspace-{workspace_id}.xlsx"
    assert sent_document.data.startswith(b"PK")
    assert (
        message.documents[0]["kwargs"]["caption"]
        == f"Workbook export for workspace {workspace_id}: topics=1, learning_items=1, topic_links=0."
    )


def test_workbook_export_handler_requires_explicit_workspace_id(tmp_path: Path) -> None:
    teacher, _ = create_teacher_workspace(tmp_path)
    message = FakeMessage(teacher)

    asyncio.run(workbook_export(message, SimpleNamespace(args="")))

    assert message.answers == [
        {"text": "Usage: /workbook_export <teacher_workspace_id>", "kwargs": {}}
    ]


def test_workbook_export_handler_rejects_student_workspace_target(tmp_path: Path) -> None:
    setup_db(tmp_path)
    teacher = make_user(1002, "Teacher")
    db.save_user(teacher)
    set_user_role(teacher.id, "teacher")
    workspace = create_workspace("Runtime", kind="student")
    add_workspace_member(workspace["workspace_id"], teacher.id, "teacher")
    message = FakeMessage(teacher)

    asyncio.run(workbook_export(message, SimpleNamespace(args=str(workspace["workspace_id"]))))

    assert message.answers == [
        {
            "text": "Workbook export is available only for teacher workspaces.",
            "kwargs": {},
        }
    ]


def test_workbook_import_usage_handler_explains_caption_flow(tmp_path: Path) -> None:
    teacher, workspace_id = create_teacher_workspace(tmp_path)
    message = FakeMessage(teacher)

    asyncio.run(workbook_import_usage(message, SimpleNamespace(args=str(workspace_id))))

    assert message.answers == [
        {
            "text": "Usage: send a .xlsx file with caption /workbook_import <teacher_workspace_id>",
            "kwargs": {},
        }
    ]


def test_workbook_import_document_handler_imports_workbook(tmp_path: Path) -> None:
    teacher, workspace_id = create_teacher_workspace(tmp_path)
    message = FakeMessage(
        teacher,
        bot=FakeBot(download_payload=build_workbook_bytes()),
        caption=f"/workbook_import {workspace_id}",
        document=make_document("import.xlsx"),
    )

    asyncio.run(workbook_import_document(message))

    assert message.answers == [
        {
            "text": (
                f"Workbook import completed for workspace {workspace_id}: "
                "created_topics=1, updated_topics=0, created_learning_items=1, "
                "updated_learning_items=0, added_topic_links=1."
            ),
            "kwargs": {},
        }
    ]


def test_workbook_import_document_handler_requires_xlsx_extension(tmp_path: Path) -> None:
    teacher, workspace_id = create_teacher_workspace(tmp_path)
    message = FakeMessage(
        teacher,
        bot=FakeBot(download_payload=b"not-used"),
        caption=f"/workbook_import {workspace_id}",
        document=make_document("import.txt"),
    )

    asyncio.run(workbook_import_document(message))

    assert message.answers == [{"text": "A .xlsx file is required.", "kwargs": {}}]


def test_workbook_import_document_handler_rejects_student_workspace_target(tmp_path: Path) -> None:
    setup_db(tmp_path)
    teacher = make_user(1003, "Teacher")
    db.save_user(teacher)
    set_user_role(teacher.id, "teacher")
    workspace = create_workspace("Runtime", kind="student")
    add_workspace_member(workspace["workspace_id"], teacher.id, "teacher")
    message = FakeMessage(
        teacher,
        bot=FakeBot(download_payload=build_workbook_bytes()),
        caption=f"/workbook_import {workspace['workspace_id']}",
        document=make_document("import.xlsx"),
    )

    asyncio.run(workbook_import_document(message))

    assert message.answers == [
        {
            "text": "Workbook import is available only for teacher workspaces.",
            "kwargs": {},
        }
    ]


def test_workbook_import_document_handler_rejects_teacher_without_membership(tmp_path: Path) -> None:
    teacher, workspace_id = create_teacher_workspace(tmp_path)
    stranger = make_user(1004, "Stranger")
    db.save_user(stranger)
    set_user_role(stranger.id, "teacher")
    message = FakeMessage(
        stranger,
        bot=FakeBot(download_payload=build_workbook_bytes()),
        caption=f"/workbook_import {workspace_id}",
        document=make_document("import.xlsx"),
    )

    asyncio.run(workbook_import_document(message))

    assert message.answers == [
        {"text": "You do not have teacher access to this workspace.", "kwargs": {}}
    ]


def test_workbook_import_document_handler_surfaces_validation_error(tmp_path: Path) -> None:
    teacher, workspace_id = create_teacher_workspace(tmp_path)
    message = FakeMessage(
        teacher,
        bot=FakeBot(download_payload=b"not-a-workbook"),
        caption=f"/workbook_import {workspace_id}",
        document=make_document("import.xlsx"),
    )

    asyncio.run(workbook_import_document(message))

    assert message.answers == [
        {"text": "Import failed: Workbook could not be opened.", "kwargs": {}}
    ]
